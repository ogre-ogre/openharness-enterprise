import paramiko

ssh = paramiko.SSHClient()
ssh.set_missing_host_key_policy(paramiko.AutoAddPolicy())
ssh.connect('10.1.7.155', username='ai-app', password='asb#1234', timeout=10)

sftp = ssh.open_sftp()

skill_md_path = '/home/ai-app/.oh-enterprise/shared/skills/docx/SKILL.md'

with sftp.open(skill_md_path) as f:
    content = f.read().decode('utf-8', errors='replace')

# 在Overview之后插入Python方案
overview_end = content.find('## Workflow Decision Tree')

python_section = '''

## ⚠️ 重要：文档生成方式变更

**由于服务器环境限制，pandoc 未安装。请使用 Python + python-docx 库直接生成Word文档。**

---

## ✅ 推荐方案：Python直接生成DOCX

使用 `python-docx` 库（已安装 v1.2.0）直接生成Word文档：

```python
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
import os

downloads_path = os.environ.get("OH_DOWNLOADS_PATH", "/tmp")
output_file = os.path.join(downloads_path, "document.docx")

doc = Document()

# 添加标题
title = doc.add_heading("文档标题", level=0)
title.alignment = WD_ALIGN_PARAGRAPH.CENTER

# 添加段落
para = doc.add_paragraph("正文内容")
para.paragraph_format.space_after = Pt(12)

# 设置字体样式
run = para.runs[0]
run.font.size = Pt(12)
run.font.name = "Arial"

# 添加粗体文本
para = doc.add_paragraph()
run = para.add_run("粗体文本")
run.bold = True

# 添加列表
doc.add_paragraph("列表项一", style="List Bullet")
doc.add_paragraph("列表项二", style="List Bullet")

# 添加表格
table = doc.add_table(rows=3, cols=2)
table.style = "Table Grid"
hdr_cells = table.rows[0].cells
hdr_cells[0].text = "列一"
hdr_cells[1].text = "列二"

doc.save(output_file)
print(f"文档已保存: {output_file}")
```

**完整功能：**
- 添加标题：`doc.add_heading(text, level)`
- 添加段落：`doc.add_paragraph(text)`
- 设置样式：`run.font.size = Pt(14)`, `run.bold = True`
- 添加表格：`doc.add_table(rows, cols)`
- 添加图片：`doc.add_picture(path, width=Inches(2))`
- 分页：`doc.add_page_break()`

---

'''

if overview_end > 0 and '⚠️ 重要' not in content:
    new_content = content[:overview_end] + python_section + content[overview_end:]
    
    with sftp.open(skill_md_path, 'w') as f:
        f.write(new_content.encode('utf-8'))
    
    print('[OK] DOCX SKILL.md updated')
else:
    print('[INFO] Already updated or marker not found')

# 验证
with sftp.open(skill_md_path) as f:
    new_content = f.read(500).decode('utf-8', errors='replace')
    if 'python-docx' in new_content:
        print('[OK] Verified: python-docx section added')

sftp.close()
ssh.close()