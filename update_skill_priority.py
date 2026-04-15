import paramiko

ssh = paramiko.SSHClient()
ssh.set_missing_host_key_policy(paramiko.AutoAddPolicy())
ssh.connect('10.1.7.155', username='ai-app', password='asb#1234', timeout=10)

sftp = ssh.open_sftp()

skill_md_path = '/home/ai-app/.oh-enterprise/shared/skills/pptx/SKILL.md'

# 读取完整内容
with sftp.open(skill_md_path) as f:
    content = f.read().decode('utf-8', errors='replace')

# 在文件开头（Overview之后）插入Python方案作为首选
overview_end = content.find('## Reading and analyzing content')

python_header = '''

## ⚠️ 重要：PPT生成方式变更

**由于服务器环境限制，请使用 Python + python-pptx 库直接生成PPT，不要使用 html2pptx.js 方案。**

Node.js 在服务器上存在 GLIBC 版本兼容问题，html2pptx.js 无法正常运行。

---

## ✅ 推荐方案：Python直接生成PPT

使用 `python-pptx` 库（已安装）直接生成PPT文件：

```python
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
import os

downloads_path = os.environ.get("OH_DOWNLOADS_PATH", "/tmp")
output_file = os.path.join(downloads_path, "presentation.pptx")

prs = Presentation()
prs.slide_width = Inches(10)
prs.slide_height = Inches(5.625)

# 添加标题幻灯片
slide = prs.slides.add_slide(prs.slide_layouts[0])
slide.shapes.title.text = "演示标题"

# 添加内容幻灯片  
slide = prs.slides.add_slide(prs.slide_layouts[1])
slide.shapes.title.text = "内容标题"

# 添加文本内容
textbox = slide.shapes.add_textbox(Inches(1), Inches(2), Inches(8), Inches(3))
tf = textbox.text_frame
tf.text = "正文内容"

prs.save(output_file)
print(f"PPT已保存: {output_file}")
```

**完整功能参考：**
- 设置字体样式：`para.font.size = Pt(24)`
- 设置颜色：`para.font.color.rgb = RGBColor(255, 100, 100)`
- 添加列表：`tf.add_paragraph()`
- 添加图片：`slide.shapes.add_picture(path, left, top, width)`
- 设置背景：`shape.fill.solid(); shape.fill.fore_color.rgb = RGBColor(26, 26, 26)`

---

'''

if overview_end > 0 and '⚠️ 重要' not in content:
    new_content = content[:overview_end] + python_header + content[overview_end:]
    
    with sftp.open(skill_md_path, 'w') as f:
        f.write(new_content.encode('utf-8'))
    
    print('[OK] SKILL.md updated - Python方案已置顶')
else:
    print('[INFO] Already updated or marker not found')

# 验证更新
with sftp.open(skill_md_path) as f:
    new_content = f.read(1000).decode('utf-8', errors='replace')
    
print('\n=== Updated content preview ===')
print(new_content[:600])

sftp.close()
ssh.close()